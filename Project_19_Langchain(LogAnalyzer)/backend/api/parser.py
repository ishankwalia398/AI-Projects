from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

SUPPORTED = {".json", ".csv", ".xls", ".xlsx", ".pdf", ".doc", ".docx", ".chls", ".chlsx", ".chlsj", ".xml", ".har"}
SENSITIVE = re.compile(r"(authorization|cookie|set-cookie|api[-_]?key|token|password|secret)", re.I)
SECRET_VALUE = re.compile(r"(?i)(\b(?:api[_-]?key|token|password|secret|authorization|cookie)\b\s*[:=]\s*)(\S+)")
BEARER_VALUE = re.compile(r"(?i)\bBearer\s+[A-Za-z0-9._~+/=-]+")
TIME = re.compile(r"\b\d{4}-\d\d-\d\d[T ]\d\d:\d\d:\d\d(?:[.,]\d+)?(?:Z|[+-]\d\d:?\d\d)?\b")
SESSION = re.compile(r"\b(?:session[-_ ]?id|x-session-id|sid)\s*[:=]\s*([\w.:-]+)", re.I)


def redact(value):
    if isinstance(value, dict):
        if isinstance(value.get("name"), str) and SENSITIVE.search(value["name"]) and "value" in value:
            value = {**value, "value": "[REDACTED]"}
        return {str(k): "[REDACTED]" if SENSITIVE.search(str(k)) else redact(v) for k, v in value.items()}
    if isinstance(value, list):
        return [redact(v) for v in value]
    if isinstance(value, str):
        return SECRET_VALUE.sub(r"\1[REDACTED]", BEARER_VALUE.sub("Bearer [REDACTED]", value))
    return value


def _headers(value):
    if isinstance(value, list):
        value = {str(h.get("name", "")): h.get("value", "") for h in value if isinstance(h, dict)}
    return redact(value if isinstance(value, dict) else {})


def _text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return redact(value)
    return json.dumps(redact(value), ensure_ascii=False, default=str)


def _entry(message="", timestamp="", session_id="", request=None, response=None, request_headers=None, response_headers=None, status=None, source=""):
    return {
        "message": _text(message).strip(),
        "raw": redact(message),
        "timestamp": _text(timestamp),
        "session_id": _text(session_id),
        "request": redact(request or {}),
        "response": redact(response or {}),
        "request_headers": _headers(request_headers),
        "response_headers": _headers(response_headers),
        "status": status,
        "source": source,
    }


def _from_har(data):
    rows = []
    for idx, item in enumerate(data.get("log", {}).get("entries", []), 1):
        req, res = item.get("request") or {}, item.get("response") or {}
        status = res.get("status")
        elapsed = item.get("time")
        session = item.get("pageref") or item.get("_sessionId") or ""
        message = " ".join(str(v) for v in [req.get("method", ""), req.get("url", ""), f"HTTP {status}" if status else "", res.get("statusText", ""), f"{elapsed} ms" if elapsed is not None else ""] if v)
        rows.append(_entry(message, item.get("startedDateTime", ""), session,
            {"method": req.get("method"), "url": req.get("url"), "body": (req.get("postData") or {}).get("text")},
            {"status": status, "status_text": res.get("statusText"), "body": (res.get("content") or {}).get("text")},
            req.get("headers"), res.get("headers"), status, f"HAR entry {idx}"))
        rows[-1]["raw"] = redact(item)
    return rows


def _from_object(data, source="JSON"):
    if isinstance(data, dict) and isinstance(data.get("log"), dict) and "entries" in data["log"]:
        return _from_har(data)
    if isinstance(data, dict):
        for key in ("entries", "logs", "events", "records", "items"):
            if isinstance(data.get(key), list):
                return _from_object(data[key], source)
        data = [data]
    if not isinstance(data, list):
        return [_entry(data, source=source)]
    rows = []
    for idx, item in enumerate(data, 1):
        if not isinstance(item, dict):
            rows.append(_entry(item, source=f"{source} {idx}"))
            continue
        # CSV exports often wrap the entire log event as JSON in a Source cell.
        # Decode containers before redaction so the agent sees fields, not escaped JSON.
        if source != "JSON":
            item = {key: _decode_container(value) for key, value in item.items()}
        lookup = {str(k).lower().replace(" ", "_"): v for k, v in item.items()}
        get = lambda *keys: next((lookup[k] for k in keys if k in lookup and lookup[k] is not None), None)
        req = get("request", "api_request") or {}
        res = get("response", "api_response") or {}
        if isinstance(req, str): req = {"body": req}
        if isinstance(res, str): res = {"body": res}
        msg = get("message", "error", "failure", "warning", "description", "detail")
        if msg is None: msg = item
        elif get("level", "severity", "type"):
            msg = f"{get('level', 'severity', 'type')}: {msg}"
        rows.append(_entry(msg, get("timestamp", "time", "datetime", "date") or "",
            get("session_id", "sessionid", "session", "sid") or "", req, res,
            get("request_headers", "req_headers") or (req.get("headers") if isinstance(req, dict) else {}),
            get("response_headers", "res_headers") or (res.get("headers") if isinstance(res, dict) else {}),
            get("status", "status_code", "http_status") or (res.get("status") if isinstance(res, dict) else None),
            f"{source} {idx}"))
        rows[-1]["raw"] = redact(item)
    return rows


def _decode_container(value):
    if isinstance(value, str) and value.lstrip().startswith(("{", "[")):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, (dict, list)):
                return parsed
        except (ValueError, RecursionError):
            pass
    return value


def _from_lines(text, source):
    rows = []
    for idx, line in enumerate(text.splitlines(), 1):
        line = line.strip()
        if not line: continue
        tm, sid = TIME.search(line), SESSION.search(line)
        rows.append(_entry(line, tm.group(0) if tm else "", sid.group(1) if sid else "", source=f"{source} line {idx}"))
    return rows


def _charles(content):
    if content.startswith(b"\xac\xed\x00\x05"):
        raise ValueError("Binary Charles .chls sessions require conversion in Charles: File → Export → XML or JSON, then upload the export.")
    try:
        data = json.loads(content.decode("utf-8-sig"))
        return _from_object(data, "Charles")
    except (UnicodeError, json.JSONDecodeError):
        pass
    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise ValueError("Charles file is neither readable JSON nor XML.") from exc
    rows = []
    for idx, node in enumerate(root.iter(), 1):
        if node.tag.lower().split("}")[-1] not in {"transaction", "entry", "session"}: continue
        fields = {child.tag.lower().split("}")[-1]: (child.text or "").strip() for child in node}
        if fields:
            rows.append(_entry(fields.get("error") or fields.get("message") or fields,
                fields.get("timestamp") or fields.get("starttime") or "",
                fields.get("sessionid") or fields.get("id") or "",
                {"method": fields.get("method"), "url": fields.get("url") or fields.get("path")},
                {"status": fields.get("status") or fields.get("responsecode")},
                status=fields.get("status") or fields.get("responsecode"), source=f"Charles {idx}"))
            rows[-1]["raw"] = redact(ET.tostring(node, encoding="unicode"))
    return rows or _from_lines(" ".join(root.itertext()), "Charles XML")


def _legacy_doc(content):
    """Best-effort text recovery from OLE Word binaries without server binaries."""
    import olefile
    if not olefile.isOleFile(io.BytesIO(content)):
        raise ValueError("This .doc is not a readable Word/OLE document.")
    with olefile.OleFileIO(io.BytesIO(content)) as ole:
        if not ole.exists("WordDocument"):
            raise ValueError("This .doc has no WordDocument stream.")
        stream = ole.openstream("WordDocument").read()
    utf16 = [m.group(0).decode("utf-16le", errors="ignore") for m in re.finditer(rb"(?:[\x20-\x7e]\x00){8,}", stream)]
    ascii_runs = [m.group(0).decode("cp1252", errors="ignore") for m in re.finditer(rb"[\x20-\x7e]{10,}", stream)]
    text = "\n".join(utf16 + ascii_runs)
    if not text.strip():
        raise ValueError("No readable text could be recovered from this legacy .doc. Save it as .docx for full extraction.")
    return _from_lines(text, "DOC")


def extract_upload(filename, content):
    """Decode file containers into source records for the LangChain parser agent."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Unsupported file type. Use JSON, CSV, XLS, XLSX, PDF, DOC, DOCX, Charles XML/JSON, or HAR.")
    if not content:
        raise ValueError("The uploaded file is empty.")
    try:
        if suffix in {".har", ".json", ".chlsj"}:
            return _from_object(json.loads(content.decode("utf-8-sig")))
        if suffix == ".csv":
            text = content.decode("utf-8-sig", errors="replace")
            return _from_object(list(csv.DictReader(io.StringIO(text))), "CSV")
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = []
            for sheet in book:
                values = sheet.values
                headers = [str(v or f"column_{i}") for i, v in enumerate(next(values, []), 1)]
                rows.extend(_from_object([dict(zip(headers, row)) for row in values], sheet.title))
            book.close()
            return rows
        if suffix == ".xls":
            import xlrd
            book = xlrd.open_workbook(file_contents=content)
            rows = []
            for sheet in book.sheets():
                if not sheet.nrows: continue
                headers = [str(v or f"column_{i}") for i, v in enumerate(sheet.row_values(0), 1)]
                rows.extend(_from_object([dict(zip(headers, sheet.row_values(i))) for i in range(1, sheet.nrows)], sheet.name))
            return rows
        if suffix == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
            if not text.strip(): raise ValueError("The PDF has no extractable text. Scanned images need OCR first.")
            return _from_lines(text, "PDF")
        if suffix == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
            return _from_lines(text, "DOCX")
        if suffix == ".doc":
            return _legacy_doc(content)
        return _charles(content)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not parse this {suffix} file: {type(exc).__name__}.") from exc


def make_markdown(filename, entries):
    lines = ["# Parsed log file", "", f"Source: {filename}", f"Generated: {datetime.now().astimezone().isoformat()}", f"Entries: {len(entries)}", ""]
    for idx, e in enumerate(entries, 1):
        if e.get("source_id"):
            lines.extend([f"Source ID: {e['source_id']}", ""])
        lines.extend([f"## Entry {idx}", "", f"- Source: {e['source']}", f"- Timestamp: {e['timestamp'] or 'Unknown'}", f"- Session ID: {e['session_id'] or 'Unknown'}", f"- Status: {e['status'] if e['status'] is not None else 'Unknown'}", "", "### Message", "", e["message"] or "(empty)", ""])
        for key, title in (("request", "API request"), ("request_headers", "Request headers"), ("response", "API response"), ("response_headers", "Response headers")):
            if e[key]: lines.extend([f"### {title}", "", "```json", json.dumps(e[key], ensure_ascii=False, indent=2, default=str), "```", ""])
    return "\n".join(lines)
