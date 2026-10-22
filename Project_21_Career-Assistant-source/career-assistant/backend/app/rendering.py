import io
import re
import zipfile
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer


def latex_escape(text):
    mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(mapping.get(c,c) for c in text)


def render(data,format):
    if format=='zip':
        buf=io.BytesIO()
        with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
            for f in ['docx','pdf','tex']:
                z.writestr('document.'+f,render(data,f))
        return buf.getvalue()
    if format=='tex':
        lines=[r'\documentclass[11pt]{article}',r'\usepackage[margin=0.8in]{geometry}',r'\usepackage[T1]{fontenc}',r'\usepackage[utf8]{inputenc}',r'\begin{document}',r'\section*{'+latex_escape(data.get('title','Career Assistant'))+'}']
        for s in data['sections']:
            lines.extend([r'\subsection*{'+latex_escape(s['heading'])+'}',latex_escape(s['text']).replace('\n',r'\par '+'\n')])
        return ('\n'.join(lines)+r'\end{document}').encode('utf-8')
    buf=io.BytesIO()
    if format=='docx':
        doc=Document()
        sec=doc.sections[0]
        sec.top_margin=sec.bottom_margin=Inches(.75)
        doc.styles['Normal'].font.name='Calibri'
        doc.styles['Normal'].font.size=Pt(10.5)
        doc.add_heading(data.get('title','Career Assistant'),0)
        for s in data['sections']:
            doc.add_heading(s['heading'],1)
            for line in s['text'].splitlines():
                doc.add_paragraph(line)
        doc.save(buf)
    elif format=='pdf':
        styles=getSampleStyleSheet()
        styles['Normal'].leading=15
        styles['Normal'].spaceAfter=7
        styles['Heading1'].fontSize=15
        styles['Heading1'].textColor=colors.HexColor('#27364c')
        # Prefer Unicode TrueType fonts on Windows/Linux when installed.
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        fonts=[Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),Path('C:/Windows/Fonts/arial.ttf')]
        font=next((p for p in fonts if p.exists()),None)
        if font:
            if 'CareerUnicode' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('CareerUnicode',str(font)))
            for name in ['Normal','Title','Heading1']:
                styles[name].fontName='CareerUnicode'
        story=[Paragraph(escape(data.get('title','Career Assistant')),styles['Title']),Spacer(1,16)]
        for s in data['sections']:
            story.append(Paragraph(escape(s['heading']),styles['Heading1']))
            for line in s['text'].splitlines():
                story.append(Paragraph(escape(line) or ' ',styles['Normal']))
        def footer(canvas,doc):
            canvas.setFont('Helvetica',8)
            canvas.setFillColor(colors.grey)
            canvas.drawRightString(550,25,str(doc.page))
        SimpleDocTemplate(buf,rightMargin=48,leftMargin=48,topMargin=42,bottomMargin=42).build(story,onFirstPage=footer,onLaterPages=footer)
    else:
        raise ValueError('Supported formats: docx, pdf, tex, zip.')
    return buf.getvalue()
