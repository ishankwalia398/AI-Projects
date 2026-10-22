import io
import json
import re
import zipfile
from datetime import date, timedelta
from unittest.mock import patch

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app import ai, rag, workflows
from app.config import settings
from app.db import File, Record, ResetToken, Session, Task, User, digest, get_record, now, put, records
from app.interviews import COMMON_FIELDS, TECH_CATEGORIES, validate_questions
from app.jobs import active_results, canonical, match, sponsorship, validate_public_url
from app.main import app
from app.rendering import render
from app.resumes import chunk, experience_months, extract
from app.security import hasher
from app.storage import cleanup, read_bytes, store, validate_upload
from conftest import login,register


DOCUMENT={'title':'Jane Example','sections':[{'heading':'Experience','text':'Engineer at ExampleCo.\nBuilt Python services. 100% & <literal>.'}]}
JOB={'company':'ExampleCo','title':'Python Engineer','location':'Bengaluru, India','country':'India','url':'https://example.com/jobs/123','description':'We need Python and SQL skills to build backend applications.','mandatory_documents':[],'mandatory_documents_known':True}


def add_job(client,**overrides):
    r=client.post('/api/jobs',json={**JOB,**overrides})
    assert r.status_code==201,r.text
    return r.json()


def resume_bytes(text='Jane Example\njane@example.com\nSkills\nPython SQL React\nExperience\nExampleCo Engineer 2020–2024'):
    d=Document()
    for line in text.splitlines():d.add_paragraph(line)
    b=io.BytesIO();d.save(b);return b.getvalue()


def ingest(client):
    r=client.post('/api/resumes',files={'file':('resume.docx',resume_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert r.status_code==200,r.text
    workflows.run_one()
    task=client.get('/api/tasks/'+r.json()['task_id']).json()
    assert task['status']=='completed',task
    rid=task['result']['resume_id']
    assert client.post('/api/resumes/'+rid+'/activate').status_code==200
    return rid


def test_registration_login_case_insensitive(env):
    c,db=env
    assert register(c,'Alice','ALICE@Example.com').status_code==201
    assert login(c,'ALICE').status_code==200
    assert c.get('/api/auth/me').json()['username']=='alice'
    with db() as session:
        user=session.query(User).one()
        assert user.password_hash.startswith('$argon2id$')
        assert hasher.verify(user.password_hash,'correct horse battery staple')


@pytest.mark.parametrize('username,email,password',[('ab','a@example.com','12345678'),('a b','a@example.com','12345678'),('alice','invalid','12345678'),('alice','a@example.com','short')])
def test_invalid_registration(env,username,email,password):
    assert register(env[0],username,email,password).status_code==422


def test_duplicate_registration(env):
    c,_=env
    register(c,'alice')
    assert register(c,'ALICE','other@example.com').status_code==409
    assert register(c,'other','ALICE@example.com').status_code==409


def test_invalid_login_generic_and_rate_limited(env):
    c,_=env
    register(c)
    a=login(c,'alice','bad').json()
    b=login(c,'missing','bad').json()
    assert a==b=={'detail':'Invalid username or password.'}
    for _ in range(8):login(c,'missing','bad')
    assert login(c,'missing','bad').status_code==429


def test_csrf_origin_cookies_logout(account):
    c,_=account
    assert c.cookies.get('career_session')
    assert 'httponly' in login(c).headers['set-cookie'].lower()
    saved=c.headers.pop('X-CSRF-Token')
    assert c.put('/api/preferences',json={}).status_code==403
    c.headers['X-CSRF-Token']=saved
    assert c.put('/api/preferences',json={},headers={'Origin':'https://evil.example'}).status_code==403
    assert c.post('/api/auth/logout').status_code==200
    assert c.get('/api/auth/me').status_code==401


def test_expired_session(account):
    c,factory=account
    with factory() as db:
        row=db.query(Session).one();row.expires_at=now()-timedelta(seconds=1);db.commit()
    assert c.get('/api/auth/me').status_code==401


def test_password_reset_single_use_and_revocation(account,monkeypatch):
    c,factory=account
    sent=[]
    class SMTP:
        def __init__(self,*a,**k):pass
        def __enter__(self):return self
        def __exit__(self,*a):pass
        def starttls(self,**k):pass
        def send_message(self,message):sent.append(message)
    monkeypatch.setattr(settings,'smtp_host','smtp.example.com')
    monkeypatch.setattr('app.main.smtplib.SMTP',SMTP)
    result=c.post('/api/auth/forgot-password',json={'email':'alice@example.com'})
    unknown=c.post('/api/auth/forgot-password',json={'email':'missing@example.com'})
    assert result.json()==unknown.json() and len(sent)==1
    token=re.search(r'\?reset=([\w-]+)',sent[0].get_content()).group(1)
    assert c.post('/api/auth/reset-password',json={'token':token,'password':'a brand new passphrase'}).status_code==200
    assert c.get('/api/auth/me').status_code==401
    assert c.post('/api/auth/reset-password',json={'token':token,'password':'another passphrase'}).status_code==400
    assert login(c,password='a brand new passphrase').status_code==200


def test_reset_invalid_expired(account):
    c,factory=account
    with factory() as db:
        user=db.query(User).one();db.add(ResetToken(id=digest('expired'),user_id=user.id,expires_at=now()-timedelta(hours=1)));db.commit()
    assert c.post('/api/auth/reset-password',json={'token':'expired','password':'valid passphrase'}).status_code==400


def test_resume_ingestion_versions_knowledge(account):
    c,factory=account
    rid=ingest(c)
    p=c.get('/api/profile').json()
    assert not p['approved'] and 'Python' in p['skills'] and p['name']==''
    p.update(name='Jane Example',approved=True,employments=[{'employer':'ExampleCo','title':'Engineer','start':'2020-01','end':'2024-01'}])
    assert c.put('/api/profile',json=p).status_code==200
    assert c.get('/api/profile').json()['experience']['months']==48
    sources=c.get('/api/knowledge?query=Python').json()
    assert sources['mode']=='lexical_only' and len(sources['sources'])==1
    ingest(c)
    assert len(c.get('/api/resumes').json())==1
    with factory() as db:
        assert records(db,db.query(User).one().id,'embedding').count()==0


def test_pdf_docx_extract_and_upload_validation():
    text=extract('resume.docx',resume_bytes())
    assert '## Skills' in text and 'Python' in text
    assert 'Python' in extract('resume.pdf',render(DOCUMENT,'pdf'))
    for name,blob in [('resume.exe',b'x'),('bad.pdf',b'not pdf'),('bad.docx',b'not zip')]:
        with pytest.raises(Exception):validate_upload(name,blob)


def test_resume_upload_invalid_and_alternate_preserves_active(account):
    c,_=account
    rid=ingest(c)
    assert c.post('/api/resumes',files={'file':('bad.pdf',b'text')}).status_code==422
    c.post('/api/resumes',files={'file':('alternate.docx',resume_bytes('Alternate candidate\nSkills\nJava Docker PostgreSQL\nExperience\nActual employer dates unknown'))})
    workflows.run_one()
    assert c.get('/api/profile').json()['resume_version_id']==rid
    assert len(c.get('/api/resumes').json())==2


def test_experience_overlap_and_incomplete():
    e=[{'start':'2020-01','end':'2022-01'},{'start':'2021-01','end':'2024-01'},{'start':'unknown','end':''}]
    assert experience_months(e,date(2025,1,1))['months']==48
    assert experience_months(e)['incomplete_roles']==1
    assert experience_months([{'start':'2024-01','end':''}],date(2025,1,1))['months']==12


def test_chunks_preserve_role_sections():
    text='## Experience\nEmployer role dates\n'+('achievement '*900)+'\n## Education\nDegree'
    parts=chunk(text)
    assert 'Employer role dates' in parts[0] and parts[1].startswith('## Education')


@pytest.mark.parametrize('text,expected',[
 ('We provide visa sponsorship.',True),('We offer relocation assistance.',True),
 ('No visa sponsorship is available.',False),('We do not offer visa sponsorship.',False),
 ('Visa sponsorship may be available.',False),('Must be eligible for relocation support.',False),
 ('We are unable to sponsor a work permit.',False)])
def test_sponsorship_conservative(text,expected):
    assert bool(sponsorship(text)['evidence']) is expected


def test_job_dedup_preserves_user_fields(account):
    c,_=account
    job=add_job(c)
    c.patch('/api/jobs/'+job['id'],json={'status':'SCREENING','comment':'My private notes'})
    duplicate=add_job(c,description=JOB['description']+' More detail.')
    assert duplicate['id']==job['id']
    assert duplicate['comment']=='My private notes' and duplicate['status']=='SCREENING'
    assert c.get('/api/jobs').json()['total']==1
    assert duplicate['history'][0]['source'].startswith('manual')


def test_canonical_urls():
    assert canonical('https://Example.com/jobs/1/?utm_source=x#part')=='https://example.com/jobs/1'


def test_visa_filter_and_domestic(account):
    c,_=account
    add_job(c)
    add_job(c,title='International Engineer',location='Berlin, Germany',country='Germany',url='https://example.com/jobs/2',description='Python engineer role. We provide visa sponsorship to successful candidates.')
    add_job(c,title='Unsupported Engineer',location='Canada',country='Canada',url='https://example.com/jobs/3')
    assert c.get('/api/jobs?module=visa').json()['total']==1
    assert c.get('/api/jobs?module=domestic').json()['total']==1


def test_matching_explained_no_invented_skills():
    m=match({'skills':['Python'],'employments':[]},{**JOB,'requirements':{'mandatory_skills':['Python','SQL']}})
    assert m['components']['mandatory_skills']['score']==15
    assert m['missing_requirements']==['SQL']
    assert 0<=m['score']<=100 and 'not an actual hiring probability' in m['label']


def test_sort_pagination_retention(account):
    c,factory=account
    with factory() as db:
        user=db.query(User).one()
        for i in range(35):
            put(db,user.id,'job',str(i),{'id':str(i),'module':'domestic','status':'NOT APPLIED','company':'Co','title':'Role','location':'India','discovered_at':(now()-timedelta(days=61 if i==0 else 1)).isoformat(),'match':{'score':i}})
        db.commit()
    first=c.get('/api/jobs').json()
    assert first['total']==34 and len(first['items'])==30
    assert first['items'][0]['match']['score']==34
    assert len(c.get('/api/jobs?page=2').json()['items'])==4
    assert c.get('/api/jobs?history=true').json()['total']==35


def test_bulk_missing_docs_comment_preserved_and_no_false_submit(account):
    c,_=account
    one=add_job(c,mandatory_documents=['Certificate'])
    two=add_job(c,title='Other role',url='https://example.com/jobs/2')
    c.patch('/api/jobs/'+one['id'],json={'comment':'Keep this note'})
    r=c.post('/api/applications/prepare',json={'job_ids':[one['id'],two['id']]}).json()
    assert r['outcomes'][0]['valid'] is False and r['outcomes'][1]['valid'] is True
    j=c.get('/api/jobs/'+one['id']).json()
    assert j['comment']=='Keep this note' and j['status']=='NOT APPLIED'
    c.post('/api/jobs/'+one['id']+'/attachments?document_type=Certificate',files={'file':('certificate.docx',resume_bytes())})
    assert c.post('/api/applications/prepare',json={'job_ids':[one['id']]}).json()['outcomes'][0]['valid']
    r=c.post('/api/applications/'+one['id']+'/confirm',json={'confirmation_reference':'Employer receipt 123','confirmed':True})
    assert r.json()['classification']=='User-provided'
    j=c.get('/api/jobs/'+one['id']).json()
    assert j['status']=='APPLIED' and j['applied_at'] and j['discovered_at']


def test_content_generation_cached_rendering_and_expiry(account,monkeypatch):
    c,factory=account
    ingest(c);p=c.get('/api/profile').json();p['approved']=True;c.put('/api/profile',json=p)
    job=add_job(c)
    calls=[]
    def generation(*args):
        calls.append(1);return {'content':DOCUMENT,'model':'test-provider-fixture','usage':{}}
    monkeypatch.setattr(ai,'generate',generation)
    for _ in range(2):
        c.post('/api/jobs/'+job['id']+'/generate');workflows.run_one()
    assert len(calls)==2  # One resume and one letter; second task reuses both.
    content=c.get('/api/content?job_id='+job['id']).json()
    assert len(content)==2
    id=content[0]['id']
    assert c.post('/api/content/'+id+'/render').status_code==409
    c.post('/api/content/'+id+'/approve')
    for format in ['pdf','docx','tex','zip']:
        r=c.post('/api/content/'+id+'/render?format='+format)
        assert r.status_code==200,r.text
        assert c.get(r.json()['url']).status_code==200
    assert len(calls)==2
    expired=c.post('/api/content/'+id+'/render?format=pdf').json()
    with factory() as db:
        f=db.get(File,expired['file_id']);f.expires_at=now()-timedelta(seconds=1);db.commit()
        assert cleanup(db)==1
        assert db.query(File).filter_by(category='original',deleted=False).count()==1
        assert records(db,db.query(User).one().id,'content').count()==2
    assert c.get(expired['url']).status_code==410
    regenerated=c.post('/api/content/'+id+'/render?format=pdf').json()
    assert regenerated['file_id']!=expired['file_id'] and len(calls)==2
    c.post('/api/auth/logout')
    assert c.get(regenerated['url']).status_code==401


def test_render_integrity_and_latex_escaping():
    assert len(PdfReader(io.BytesIO(render(DOCUMENT,'pdf'))).pages)>=1
    doc=Document(io.BytesIO(render(DOCUMENT,'docx')))
    assert any('Python' in p.text for p in doc.paragraphs)
    assert b'100\\% \\&' in render(DOCUMENT,'tex')
    with zipfile.ZipFile(io.BytesIO(render(DOCUMENT,'zip'))) as z:
        assert set(z.namelist())=={'document.pdf','document.docx','document.tex'}


def test_cleanup_only_temporary_preserves_attachment(account):
    c,factory=account
    job=add_job(c)
    c.post('/api/jobs/'+job['id']+'/attachments?document_type=Resume',files={'file':('resume.docx',resume_bytes())})
    with factory() as db:
        user=db.query(User).one()
        original=store(db,user.id,'original.pdf',b'%PDF-original','original')
        temp=store(db,user.id,'temp.pdf',b'%PDF-temp','temporary',expires_at=now()-timedelta(hours=1))
        db.commit();assert cleanup(db)==1
        assert read_bytes(original.key)==b'%PDF-original'
        assert db.query(File).filter_by(category='attachment',deleted=False).count()==1


def test_cross_user_isolation(account):
    a,factory=account
    job=add_job(a)
    other=TestClient(app)
    register(other,'bob');login(other,'bob')
    assert other.get('/api/jobs').json()['total']==0
    assert other.get('/api/jobs/'+job['id']).status_code==404
    assert other.patch('/api/jobs/'+job['id'],json={'comment':'Intrusion'}).status_code==404
    assert other.post('/api/jobs/'+job['id']+'/generate').status_code==404
    assert other.post('/api/applications/prepare',json={'job_ids':[job['id']]}).status_code==404
    other.close()


def test_unconfigured_ai_fails_truthfully(account):
    c,_=account
    ingest(c);p=c.get('/api/profile').json();p['approved']=True;c.put('/api/profile',json=p)
    j=add_job(c)
    result=c.post('/api/jobs/'+j['id']+'/generate').json();workflows.run_one()
    task=c.get('/api/tasks/'+result['task_id']).json()
    assert task['status']=='failed' and 'not configured' in task['error']
    assert c.get('/api/content').json()==[]


def test_interview_question_validator_requires_answers():
    q={k:'detail' for k in COMMON_FIELDS}
    q.update(id='T1',category=TECH_CATEGORIES[0],source_classification='General Preparation',follow_ups=[{'question':'Why?','answer':'Because.'}])
    qs=[{**q,'id':str(i),'question':f'Question {i}?'} for i in range(3)]
    validate_questions({'questions':qs},TECH_CATEGORIES[0],'technical')
    qs[0]['detailed_answer']=''
    with pytest.raises(ValueError):validate_questions({'questions':qs},TECH_CATEGORIES[0],'technical')


def test_rag_filters_always_tenant_namespace_version():
    filt=rag.tenant_filter('user-1','candidate','version-2').model_dump()
    assert {r['key'] for r in filt['must']}=={'user_id','namespace','version'}


@pytest.mark.parametrize('url',['http://example.com','https://127.0.0.1','https://localhost','https://example.com:8000','https://user:secret@example.com'])
def test_ssrf_urls_rejected(url):
    with pytest.raises(ValueError):validate_public_url(url)


def test_account_delete_preserves_other_user(account):
    c,factory=account
    add_job(c);register(c,'bob')
    assert c.request('DELETE','/api/account',json={'password':'wrong password'}).status_code==403
    assert c.request('DELETE','/api/account',json={'password':'correct horse battery staple'}).status_code==200
    assert c.get('/api/auth/me').status_code==401
    with factory() as db:
        assert db.query(User).one().username=='bob'
        assert db.query(Record).filter_by(kind='job').count()==0
