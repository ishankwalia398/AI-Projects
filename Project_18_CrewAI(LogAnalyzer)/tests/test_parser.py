import io
import csv
import json
import unittest

from backend.api.parser import make_markdown, extract_upload, redact


class ParserTests(unittest.TestCase):
    def test_csv_embedded_json_is_decoded_and_redacted(self):
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(["Timestamp", "Source", "Application", "Subsystem"])
        writer.writerow(["2026-09-22T10:00:00Z", json.dumps({"message": "Request failed", "severity": "ERROR", "token": "synthetic-secret"}), "api", "test"])
        records = extract_upload("export.csv", stream.getvalue().encode())
        self.assertEqual(records[0]["raw"]["Source"]["message"], "Request failed")
        self.assertEqual(records[0]["raw"]["Source"]["token"], "[REDACTED]")
        self.assertNotIn("synthetic-secret", json.dumps(records))

    def test_har_preserves_http_context_and_redacts_secrets(self):
        har = {"log": {"entries": [{
            "startedDateTime": "2026-09-21T08:00:00Z", "pageref": "session-42", "time": 3000,
            "request": {"method": "GET", "url": "https://api.test/orders", "headers": [{"name": "Authorization", "value": "Bearer secret"}]},
            "response": {"status": 503, "statusText": "Unavailable", "headers": [{"name": "Retry-After", "value": "30"}], "content": {"text": "upstream error"}},
        }]}}
        entries = extract_upload("traffic.har", json.dumps(har).encode())
        findings = entries
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["session_id"], "session-42")
        self.assertEqual(findings[0]["request_headers"]["Authorization"], "[REDACTED]")
        self.assertEqual(findings[0]["response_headers"]["Retry-After"], "30")
        self.assertNotIn("Bearer secret", make_markdown("traffic.har", entries))
        self.assertNotIn("Bearer secret", json.dumps(entries[0]["raw"]))

    def test_csv_levels_and_timeout(self):
        csv = b"timestamp,session_id,level,message\n2026-09-21T09:00:00Z,s1,warning,Retrying request\n2026-09-21T09:01:00Z,s1,error,Request timed out\n"
        findings = extract_upload("events.csv", csv)
        self.assertEqual([f["raw"]["level"] for f in findings], ["warning", "error"])

    def test_xlsx_rows(self):
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active
        sheet.append(["timestamp", "session_id", "message"])
        sheet.append(["2026-09-21T09:00:00Z", "s2", "FAILURE checkout rejected"])
        stream = io.BytesIO(); book.save(stream)
        findings = extract_upload("events.xlsx", stream.getvalue())
        self.assertEqual(findings[0]["raw"]["message"], "FAILURE checkout rejected")

    def test_docx_lines(self):
        from docx import Document
        doc = Document(); doc.add_paragraph("2026-09-21T10:00:00Z session_id=s3 WARNING queue degraded")
        stream = io.BytesIO(); doc.save(stream)
        findings = extract_upload("events.docx", stream.getvalue())
        self.assertEqual(findings[0]["session_id"], "s3")

    def test_unknown_extension_rejected(self):
        with self.assertRaises(ValueError): extract_upload("bad.exe", b"hello")

    def test_charles_json_and_binary_message(self):
        rows = extract_upload("traffic.chlsj", b'[{"message":"WARNING slow request"}]')
        self.assertEqual(rows[0]["raw"]["message"], "WARNING slow request")
        with self.assertRaisesRegex(ValueError, "Binary Charles"):
            extract_upload("traffic.chls", b"\xac\xed\x00\x05anything")

    def test_inline_bearer_redaction(self):
        self.assertNotIn("very-secret", redact("Authorization: Bearer very-secret"))


if __name__ == "__main__":
    unittest.main()
